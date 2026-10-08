# -*- coding: utf-8 -*-
import sys, os, glob, io
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.ns import qn

def iter_block_items(parent):
    body = parent.element.body
    for child in body.iterchildren():
        if child.tag == qn('w:p'):
            yield Paragraph(child, parent)
        elif child.tag == qn('w:tbl'):
            yield Table(child, parent)

def dump(path):
    out = []
    out.append('=' * 100)
    out.append('FILE: ' + os.path.basename(path))
    out.append('=' * 100)
    doc = Document(path)
    ti = 0
    for block in iter_block_items(doc):
        if isinstance(block, Paragraph):
            t = block.text.strip()
            if t:
                style = block.style.name if block.style is not None else ''
                out.append(f'[P:{style}] {t}')
        else:
            ti += 1
            out.append(f'--- TABLE {ti} ({len(block.rows)}x{len(block.columns)}) ---')
            for r in block.rows:
                cells = []
                seen = set()
                for c in r.cells:
                    if id(c._tc) in seen:
                        continue
                    seen.add(id(c._tc))
                    cells.append(c.text.strip().replace('\n', ' / '))
                out.append('  | ' + ' | '.join(cells) + ' |')
            out.append('--- END TABLE ---')
    # headers/footers
    for si, sec in enumerate(doc.sections):
        hdr = ' / '.join(p.text.strip() for p in sec.header.paragraphs if p.text.strip())
        ftr = ' / '.join(p.text.strip() for p in sec.footer.paragraphs if p.text.strip())
        if hdr or ftr:
            out.append(f'[SEC{si} HEADER] {hdr}')
            out.append(f'[SEC{si} FOOTER] {ftr}')
    return '\n'.join(out)

if __name__ == '__main__':
    folder = sys.argv[1]
    files = sorted(glob.glob(os.path.join(folder, '*.docx')))
    allout = []
    for f in files:
        try:
            allout.append(dump(f))
        except Exception as e:
            allout.append(f'ERROR {f}: {e}')
    txt = '\n\n'.join(allout)
    with io.open(sys.argv[2], 'w', encoding='utf-8') as fh:
        fh.write(txt)
    print('written', sys.argv[2], len(txt))
